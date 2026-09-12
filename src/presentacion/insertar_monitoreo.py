# -*- coding: utf-8 -*-
"""Inserta una lámina sobre el monitoreo de distribuciones registrado en MLflow, después de la 13,
y renumera las siguientes. Conserva todo lo demás del archivo del candidato."""
# Se ejecutan SOBRE el archivo ya generado, para conservar los ajustes hechos a mano en PowerPoint.
# No vuelvas a correr construir.py después de editar el pptx: lo sobrescribe.
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
ARCHIVO = ROOT / "outputs" / "presentacion_opciones_de_pago.pptx"
FIG = ROOT / "outputs" / "ppt"

AZUL = RGBColor(0x2A, 0x78, 0xD6)
VERDE = RGBColor(0x2E, 0x9E, 0x6B)
GRIS = RGBColor(0x6B, 0x6B, 0x6B)
TEXTO = RGBColor(0x33, 0x33, 0x33)
CLARO = RGBColor(0xF3, 0xF5, 0xF8)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)


def caja(slide, x, y, w, h, texto, tam=15, negrita=False, color=TEXTO, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.0
        p.space_after = Pt(3) if linea.strip() else Pt(1)
        t, neg = linea, negrita
        if t.startswith("**") and t.endswith("**"):
            t, neg = t[2:-2], True
        r = p.add_run()
        r.text = t
        r.font.size = Pt(tam)
        r.font.bold = neg
        r.font.color.rgb = color
        r.font.name = "Segoe UI"
    return tb


def bloque(slide, x, y, w, h, relleno):
    s = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = relleno
    s.line.fill.background(); s.shadow.inherit = False
    return s


prs = Presentation(str(ARCHIVO))
print(f"antes: {len(prs.slides)} diapositivas")

s = prs.slides.add_slide(prs.slide_layouts[6])
banda = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.14))
banda.fill.solid(); banda.fill.fore_color.rgb = AZUL
banda.line.fill.background(); banda.shadow.inherit = False

caja(s, 0.7, 0.45, 12, 0.9, "El modelo avisa cuando los datos cambian", tam=32, negrita=True, color=AZUL)
caja(s, 0.7, 1.32, 12, 0.5, "Un trabajo mensual mide la deriva de cada variable y la registra en MLflow", tam=15, color=GRIS)

s.shapes.add_picture(str(FIG / "10_monitoreo.png"), Inches(0.6), Inches(1.95), width=Inches(8.1))

bloque(s, 9.0, 2.0, 3.7, 2.15, CLARO)
caja(s, 9.3, 2.25, 3.1, 1.8, "**Qué mide**\nCompara la distribución de\ncada variable en enero contra\nla del entrenamiento.\nTres estados: estable, aviso\ny alarma.", tam=13.5)

bloque(s, 9.0, 4.35, 3.7, 2.15, VERDE)
caja(s, 9.3, 4.6, 3.1, 1.8, "**Queda en MLflow**\nCada corrida guarda el\nindicador por variable, el\nreparto por estado y las\ngráficas, con la fecha.\nAsí se compara mes a mes.", tam=13.5, color=BLANCO)

bloque(s, 0.6, 5.5, 8.1, 1.45, CLARO)
caja(s, 0.95, 5.72, 7.6, 1.1,
     "**La lectura: varias variables se mueven, pero la predicción no.**\n"
     "Las tres tasas históricas se disparan porque en agosto son nulas por construcción. Aun así, la salida\n"
     "del modelo se mantiene estable y coherente con el resultado real de enero.", tam=13)

caja(s, 12.4, 6.95, 0.6, 0.3, "0", tam=11, color=GRIS, align=PP_ALIGN.RIGHT)

NOTAS = (
    "Esta lámina cierra la parte de puesta en producción, y responde a la pregunta de qué pasa cuando el mundo cambia.\n"
    "Un modelo entrenado con datos de 2023 no sirve para siempre. Monté un trabajo mensual que compara la distribución de "
    "cada una de las 100 variables en el mes nuevo contra la que tenían en el entrenamiento, y las clasifica en tres estados: "
    "estable, aviso y alarma.\n"
    "A la izquierda, el resultado real de enero: 71 variables estables, 15 en aviso y 14 en alarma.\n"
    "A la derecha, las que más se movieron. Y aquí viene lo importante, porque un número alto no siempre es un problema. "
    "Las tres primeras son tasas históricas de aceptación por grupo, y se disparan por construcción: en agosto de 2023 no había "
    "historia previa, así que eran nulas, y en enero ya acumulan cinco meses. No es que el mundo haya cambiado, es que la variable "
    "necesita tiempo para existir. Otras, como el comportamiento de pago previo, sí parecen estacionalidad de fin de año.\n"
    "El dato que me tranquiliza es el de abajo: aunque varias variables se movieron, la distribución de la predicción se mantuvo "
    "estable, y eso es coherente con el resultado que obtuve en enero.\n"
    "Todo queda registrado en MLflow con su fecha, así que el mes siguiente puedo comparar y ver si la deriva crece. "
    "Y hay una regla escrita: si el desempeño cae por debajo del umbral dos meses seguidos, se dispara el reentrenamiento.")

def poner_notas(dia, texto):
    """Escribe las notas del orador aunque la plantilla no traiga marcador de cuerpo."""
    ns = dia.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        tf = next((sh.text_frame for sh in ns.shapes if sh.has_text_frame), None)
    if tf is None:   # la plantilla no trae marcador de notas: se copia el de otra diapositiva
        import copy
        modelo = next((d for d in prs.slides if d is not dia and d.has_notes_slide and d.notes_slide.notes_text_frame is not None), None)
        if modelo is None:
            print("  aviso: no se pudieron escribir las notas de esta lamina")
            return
        ph = modelo.notes_slide.notes_placeholder
        ns.shapes._spTree.append(copy.deepcopy(ph._element))
        tf = ns.notes_text_frame
    tf.text = texto


poner_notas(s, NOTAS)

# posicion: justo despues de "El modelo esta en produccion"
destino = None
for i, dia in enumerate(prs.slides):
    if any(sh.has_text_frame and sh.text_frame.text.startswith("El modelo está en producción") for sh in dia.shapes):
        destino = i + 1
        break
if destino is None:
    raise SystemExit("no encontre la lamina de produccion")
print("se insertara en la posicion", destino + 1)

lst = prs.slides._sldIdLst
ids = list(lst)
nueva = ids[-1]
lst.remove(nueva)
lst.insert(destino, nueva)
print("insertada")

for i, dia in enumerate(prs.slides, 1):
    for sh in dia.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.isdigit() and abs(sh.left / 914400 - 12.4) < 0.2 and sh.top / 914400 > 6.5:
            if t != str(i) and sh.text_frame.paragraphs[0].runs:
                sh.text_frame.paragraphs[0].runs[0].text = str(i)
print("renumeradas")

prs.save(str(ARCHIVO))
print(f"despues: {len(prs.slides)} diapositivas")
